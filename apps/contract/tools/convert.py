import openpyxl, json, re, sys, datetime, random, string, collections
src=sys.argv[1]; out=sys.argv[2]
wb=openpyxl.load_workbook(src,read_only=True,data_only=True)
def rid(p): return p+''.join(random.choices(string.ascii_lowercase+string.digits,k=10))
NOW=datetime.datetime.now().isoformat(timespec='seconds')
def s(v):
    if v is None: return ''
    if isinstance(v,datetime.datetime): return v.strftime('%Y-%m-%d')
    v=str(v).replace('　',' ').strip()
    return v
def num(v):
    if v is None or v=='': return 0
    if isinstance(v,(int,float)): return round(v)
    v=str(v).replace(',','').strip()
    try: return round(float(v))
    except: return 0
DH=re.compile(r'^\d{3}-\d{3,4}$')
def norm_dh(v):
    v=s(v); cancel=False
    if v[:1] in 'aA': v=v[1:]; cancel=True
    if v[-1:] in 'aA': v=v[:-1]; cancel=True
    return v, cancel
STAGES=['계약금','중도금 1차','중도금 2차','중도금 3차','중도금 4차','중도금 5차','잔금']
# units
units=[]
for row in wb['분양가테이블'].iter_rows(min_row=3,max_col=35,values_only=True):
    if not row[2]: continue
    am={st:num(row[25+i]) for i,st in enumerate(STAGES)}
    units.append(dict(id=s(row[2]),dong=s(row[3]),ho=s(row[4]),floor=num(row[5]),line=num(row[6]),type=s(row[7]),
      exArea=row[8] or 0,resArea=row[10] or 0,supArea=row[12] or 0,etcArea=row[14] or 0,conArea=row[16] or 0,landArea=row[18] or 0,
      priceLand=num(row[20]),priceBuild=num(row[21]),vat=num(row[22]),price=num(row[23]),amounts=am,updatedAt=NOW))
uset={u['id'] for u in units}
# contracts
contracts=[]
for row in wb['계약관리'].iter_rows(min_row=5,max_col=27,values_only=True):
    if row[9] is None: continue
    dh,cancel=norm_dh(row[9])
    if dh=='대체계약': dh=''; cancel=True
    status='해약' if cancel else '정상'
    memo2=' / '.join(s(x) for x in row[24:27] if s(x))
    contracts.append(dict(id=rid('c'),dongho=dh,status=status,report=s(row[1]) or '미보고',ctype=s(row[7]),date=s(row[8]),
      name=s(row[10]),rrn=s(row[11]),addr=s(row[12]),mail=s(row[13]),phone=s(row[14]),email=s(row[15]),
      stype=(s(row[16]).replace('일반공급','일반').replace('신혼','신혼부부').replace('신혼부부부부','신혼부부')),fee=s(row[17]),team=s(row[18]),staff=s(row[19]),
      note=s(row[22]),memo=s(row[23]),memo2=memo2,updatedAt=NOW))
active={c['dongho'] for c in contracts if c['status']=='정상'}
def simple(sheet, minrow, cols, kind):
    res=[]
    for row in wb[sheet].iter_rows(min_row=minrow,max_col=16,values_only=True):
        d=dict((k,row[i]) for k,i in cols.items())
        if not s(d.get('name')) and not s(d.get('dongho')): continue
        if s(d.get('dongho'))=='끝': continue
        dh,cancel=norm_dh(d['dongho'])
        st=s(d.get('status'))
        if st not in ('전환','취소','환불'):
            st = ('전환' if dh in active else '취소') if cancel else '진행'
        r=dict(id=rid(kind),date=s(d['date']),dongho=dh if DH.match(dh) else '',status=st,name=s(d['name']),phone=s(d.get('phone')),
               plan=s(d.get('plan')),team=s(d.get('team')),staff=s(d.get('staff')),note=s(d.get('note')),updatedAt=NOW)
        if not DH.match(dh) and dh: r['note']=(r['note']+' (원동호:'+dh+')').strip()
        res.append(r)
    return res
pre=simple('가계약관리',5,dict(date=5,dongho=6,status=7,name=8,phone=9,plan=11,team=12,staff=13,note=14),'p')
resv=[]
empty=0
rows=[]
for row in wb['동호지정관리'].iter_rows(min_row=5,max_col=16,values_only=True):
    if row[7] is None and row[5] is None:
        empty+=1
        if empty>300: break
        continue
    empty=0; rows.append(row)
for row in rows:
    dh,cancel=norm_dh(row[5]); st=s(row[6])
    if st not in ('전환','취소','환불'):
        st=('전환' if dh in active else '취소') if cancel else '진행'
    resv.append(dict(id=rid('r'),date=s(row[4]),dongho=dh,status=st,name=s(row[7]),phone=s(row[8]),plan=s(row[9]),team=s(row[10]),staff=s(row[11]),note=s(row[12]),updatedAt=NOW))
locks=[]
for row in wb['잠금동호관리'].iter_rows(min_row=3,max_col=4,values_only=True):
    if not row[1]: continue
    dh,cancel=norm_dh(row[1])
    note=s(row[3]); reason=s(row[2])
    if reason in ('해제',) or '해제' in reason: cancel=True
    locks.append(dict(id=rid('l'),date=s(row[0]),dongho=dh,reason=reason,note=note,status='해제' if cancel else '잠금',updatedAt=NOW))
# deposits
SPECIAL={'대체지급','예금결산','예금결산이자','입금오류','취소','해약','착오입금'}
deps=[]; seen=set(); dupcnt=0
for row in wb['수금현황'].iter_rows(min_row=10,max_row=6000,max_col=14,values_only=True):
    if row[1] is None: continue
    t=s(row[1]); outv=num(row[4]); inv=num(row[5]); bal=num(row[6])
    key=f'{t}|{inv}|{outv}|{bal}'
    if key in seen: dupcnt+=1; continue
    seen.add(key)
    cat=s(row[9]); k=s(row[10]); dh,cancel=norm_dh(k)
    note=''
    if not DH.match(dh):
        if k=='환불신청': cat=cat or '환불'
        elif k in SPECIAL: 
            cat = {'예금결산이자':'예금결산','착오입금':'환불','해약':cat or '환불','취소':'환불'}.get(k,k) if not cat or cat in ('환불',) and k in ('착오입금',) else cat
            if k in ('취소','해약','착오입금'): note=k
        elif k: note=k
        dh=''
    if cat=='환불' and outv>0: cat='환불지급'
    deps.append(dict(id=rid('d'),k=key,t=t,jeok=s(row[2]),desc=s(row[3]),out=outv,inn=inv,bal=bal,branch=s(row[7]),cat=cat,dongho=dh,name=s(row[11]),note=note,updatedAt=NOW))
def tsort(d): return d['t']
deps.sort(key=tsort)  # stable
settings=dict(project=dict(name='양주역 중흥S-클래스',lessor='중흥토건㈜',bizNo='416-81-13368',bank='우리은행',account='1005604791971',moneyLabel='임대보증금'),
  stages=STAGES,
  special={'59C':{'청년':12,'신혼부부':42,'고령자':6},'59D':{'청년':2,'신혼부부':11,'고령자':1},'84C':{'청년':7,'신혼부부':26,'고령자':3},'84D':{'청년':2,'신혼부부':11,'고령자':1}},
  depositCats=STAGES+['환불','환불지급','대체지급','예금결산','입금오류','기타'])
data=dict(version=1,exportedAt=NOW,source='excel:'+src.split('/')[-1],settings=settings,
  units=units,contracts=contracts,pre=pre,resv=resv,locks=locks,deposits=deps)
json.dump(data,open(out,'w',encoding='utf-8'),ensure_ascii=False)
print('units',len(units),'contracts',len(contracts),'active',sum(c['status']=='정상' for c in contracts),'pre',len(pre),'resv',len(resv),'locks',len(locks),'deps',len(deps),'dup',dupcnt)
print(collections.Counter(d['cat'] for d in deps))
print(collections.Counter(p['status'] for p in pre), collections.Counter(p['status'] for p in resv), collections.Counter(p['status'] for p in locks))
print(collections.Counter(c['stype'] for c in contracts))
